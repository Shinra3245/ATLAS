"""Inspect all received files without changing incoming; XLSX profiling via openpyxl."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

ROOT = Path(__file__).resolve().parents[2]


def sha256(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def tables(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in book:
            rows = [list(map(scalar, row)) for row in sheet.iter_rows(values_only=True)]
            # Remove only completely empty trailing rows and columns in the profile.
            while rows and all(v is None for v in rows[-1]):
                rows.pop()
            if rows:
                width = max((i + 1 for row in rows for i, v in enumerate(row) if v is not None), default=0)
                rows = [row[:width] for row in rows]
            yield sheet.title, rows
    finally:
        book.close()


def inventory():
    profiles = []
    for path in sorted((ROOT / 'data/incoming').rglob('*')):
        if not path.is_file() or path.name == '.gitkeep':
            continue
        profile = dict(path=str(path.relative_to(ROOT)), name=path.name,
                       format=path.suffix.lower(), bytes=path.stat().st_size,
                       sha256=sha256(path), sheets=[], error=None)
        if path.suffix.lower() == '.xlsx':
            try:
                book = openpyxl.load_workbook(path, read_only=True, data_only=False)
                # openpyxl supplies default timestamps when docProps/core.xml is absent.
                # Read XML directly so those defaults are never mistaken for source dates.
                with zipfile.ZipFile(path) as archive:
                    profile['workbook_metadata'] = {}
                    if 'docProps/core.xml' in archive.namelist():
                        for elem in ET.fromstring(archive.read('docProps/core.xml')):
                            profile['workbook_metadata'][elem.tag.split('}')[-1]] = elem.text
                profile['formulas'] = {sheet.title: sum(isinstance(c.value, str) and c.value.startswith('=')
                                                      for row in sheet for c in row) for sheet in book}
                book.close()
                for title, rows in tables(path):
                    header, body = (rows[0], rows[1:]) if rows else ([], [])
                    columns = []
                    for i, name in enumerate(header):
                        vals = [row[i] if i < len(row) else None for row in body]
                        present = [v for v in vals if v is not None]
                        numeric = [v for v in present if isinstance(v, (int, float)) and not isinstance(v, bool)]
                        counts = Counter(str(v) for v in present)
                        columns.append(dict(name=name, nulls=len(vals)-len(present),
                                            types=dict(Counter(type(v).__name__ for v in present)),
                                            distinct=len(counts), examples=list(counts)[:6],
                                            top=counts.most_common(8), min=min(numeric) if numeric else None,
                                            max=max(numeric) if numeric else None,
                                            observed_date_min=min(map(str,present)) if str(name).upper() in ['PERIODO','FECHA_ACT','ANIO','AÑO'] and present else None,
                                            observed_date_max=max(map(str,present)) if str(name).upper() in ['PERIODO','FECHA_ACT','ANIO','AÑO'] and present else None))
                    body_hashes = [json.dumps(row, ensure_ascii=False, sort_keys=True) for row in body]
                    semantic_hash = hashlib.sha256(json.dumps(rows, ensure_ascii=False).encode()).hexdigest()
                    profile['sheets'].append(dict(name=title, records=len(body), columns=len(header),
                                                 fields=columns, samples=body[:3],
                                                 duplicate_rows=len(body_hashes)-len(set(body_hashes)),
                                                 semantic_sha256=semantic_hash,
                                                 coordinate_columns=[str(n) for n in header if any(k in str(n).upper() for k in ['LONGITUD','LATITUD','LAT_REP','LON_REP','BBOX_','EPSG','UTMX','UTMY','LCCX','LCCY'])],
                                                 municipality_counts=dict(Counter(str(row[header.index('NOM_MUN')]) for row in body)) if 'NOM_MUN' in header else {},
                                                 metadata_rows=rows if len(rows) <= 120 else None))
            except Exception as exc:
                profile['error'] = f'{type(exc).__name__}: {exc}'
        else:
            profile['error'] = 'FORMAT_PENDING_INSPECTION'
        profiles.append(profile)
    return profiles


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, default=ROOT / 'data/metadata/inventory.json')
    args = parser.parse_args()
    profiles = inventory()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(profiles, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    for p in profiles:
        print(p['path'], p['bytes'], p['sha256'], p['error'] or '')
        for s in p['sheets']:
            print(' ', s['name'], s['records'], 'rows', s['columns'], 'columns')


if __name__ == '__main__':
    main()
